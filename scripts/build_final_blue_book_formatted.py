import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def clean_xml(text):
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', str(text))

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def delete_paragraph(paragraph):
    p = paragraph._p
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)

def main():
    base_dir = r'C:\Users\Sumit\OneDrive\Desktop\FraudShield_AI\FraudShield_AI'
    template_path = os.path.join(base_dir, 'docs', 'Draft Blue book format for 2026-27 -Final.docx')
    output_path = os.path.join(base_dir, 'docs', 'FraudShield_AI_Final_Blue_Book.docx')
    markdown_source = os.path.join(base_dir, 'docs', 'FraudShield_AI_Blue_Book_Report.md')
    figures_dir = os.path.join(base_dir, 'docs', 'figures')

    print("Loading official TCET template...")
    doc = docx.Document(template_path)

    # 1. Update Title Page
    p_title = doc.paragraphs[4]
    p_title.text = ""
    run_t = p_title.add_run("FraudShield AI: An Intelligent Real-Time Financial Transaction & Credit Card Fraud Detection Platform Using Hybrid Stacking Ensemble and Explainable AI (XAI)")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.15

    # Clear instructions <Times New Roman...>
    for p_idx in [3, 5, 7, 9, 11, 15, 18, 20, 23, 26]:
        if p_idx < len(doc.paragraphs):
            doc.paragraphs[p_idx].text = ""

    # Paragraph 13: "Submitted by"
    p_sub_by = doc.paragraphs[13]
    p_sub_by.text = ""
    r_sb = p_sub_by.add_run("Submitted by")
    r_sb.font.name = "Times New Roman"
    r_sb.font.size = Pt(14)
    r_sb.font.bold = True
    p_sub_by.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 14: Students with correct roll numbers: 10, 21, 19
    p_stud = doc.paragraphs[14]
    p_stud.text = ""
    for name in ["Ashmit Singh (Roll No. 10)", "Sumit Singh (Roll No. 21)", "Shivam Singh (Roll No. 19)"]:
        r = p_stud.add_run(name + "\n")
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
    p_stud.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 17: "Under the Guidance of"
    p_guide_lbl = doc.paragraphs[17]
    p_guide_lbl.text = ""
    r_gl = p_guide_lbl.add_run("Under the Guidance of")
    r_gl.font.name = "Times New Roman"
    r_gl.font.size = Pt(12)
    r_gl.font.bold = True
    r_gl.font.italic = True
    p_guide_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 19: Guide Name
    p_guide = doc.paragraphs[19]
    p_guide.text = ""
    r_g = p_guide.add_run("Ms. Tanmayi Nagale")
    r_g.font.name = "Times New Roman"
    r_g.font.size = Pt(14)
    r_g.font.bold = True
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 22: Designation
    p_desig = doc.paragraphs[22]
    p_desig.text = ""
    r_d = p_desig.add_run("Assistant Professor")
    r_d.font.name = "Times New Roman"
    r_d.font.size = Pt(14)
    r_d.font.bold = True
    p_desig.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 25: Dept
    p_dept = doc.paragraphs[25]
    p_dept.text = ""
    r_dp = p_dept.add_run("Department of Computer Engineering\n(Academic Year 2026-27)")
    r_dp.font.name = "Times New Roman"
    r_dp.font.size = Pt(14)
    r_dp.font.bold = True
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2. Update Certificate (Paragraph 35 / 36)
    for p_idx in [34, 35, 36]:
        if p_idx < len(doc.paragraphs):
            txt = doc.paragraphs[p_idx].text
            if "certify that" in txt.lower() or "bonafide" in txt.lower() or "<times new roman" in txt.lower():
                if "<times new roman" in txt.lower():
                    doc.paragraphs[p_idx].text = ""
                else:
                    p_cert = doc.paragraphs[p_idx]
                    p_cert.text = ""
                    r_c = p_cert.add_run('This is to certify that the project entitled “FraudShield AI: An Intelligent Real-Time Financial Transaction & Credit Card Fraud Detection Platform Using Hybrid Stacking Ensemble and Explainable AI (XAI)” is a bonafide work of Ashmit Singh (Roll No. 10), Sumit Singh (Roll No. 21), and Shivam Singh (Roll No. 19) submitted to the Thakur College of Engineering and Technology, Mumbai (An Autonomous College affiliated to University of Mumbai) in partial fulfillment of the requirement for the Project-I for award of the degree of “Bachelor of Engineering” in “Computer Engineering”.')
                    r_c.font.name = "Times New Roman"
                    r_c.font.size = Pt(12)
                    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p_cert.paragraph_format.line_spacing = 1.5

    # Table 0 Signatures
    t0 = doc.tables[0]
    t0.rows[0].cells[0].text = "Signature with Date: -----------------\n\nName of Guide: Ms. Tanmayi Nagale\nDesignation: Assistant Professor"
    t0.rows[0].cells[1].text = "Signature with Date: --------------------\n\nName of HOD: Dr. Vaishali Kaiche\nName of Department: Department of Computer Engineering"
    for r in t0.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)

    # 3. Update Approval Certificate
    for p_idx in range(45, 55):
        if p_idx < len(doc.paragraphs):
            txt = doc.paragraphs[p_idx].text
            if "approved for" in txt.lower() or "report entitled" in txt.lower():
                p_app = doc.paragraphs[p_idx]
                p_app.text = ""
                r_a = p_app.add_run('This project report entitled “FraudShield AI: An Intelligent Real-Time Financial Transaction & Credit Card Fraud Detection Platform Using Hybrid Stacking Ensemble and Explainable AI (XAI)” by Ashmit Singh (Roll No. 10), Sumit Singh (Roll No. 21), and Shivam Singh (Roll No. 19) is approved for the degree of “Bachelor of Engineering” in “Computer Engineering”.')
                r_a.font.name = "Times New Roman"
                r_a.font.size = Pt(12)
                p_app.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_app.paragraph_format.line_spacing = 1.5
            elif "<times new roman" in txt.lower():
                doc.paragraphs[p_idx].text = ""

    # 4. Update Acknowledgement & Remove unwanted template draft notes
    for p_idx in range(60, 90):
        if p_idx < len(doc.paragraphs):
            p = doc.paragraphs[p_idx]
            txt = p.text
            if "unfair if i do not acknowledge" in txt.lower() or "stick to our backs" in txt.lower() or "<times new roman" in txt.lower() or "<line spacing" in txt.lower():
                p.text = ""

    # Set clean formal Acknowledgement
    p_ack1 = doc.paragraphs[63]
    p_ack1.text = "We express our deep sense of gratitude to our project guide Ms. Tanmayi Nagale, Assistant Professor, Department of Computer Engineering, for her invaluable guidance, constant support, and insightful encouragement throughout the course of this work. We also extend our sincere appreciation to Dr. Vaishali Kaiche, Head of the Department of Computer Engineering, and Dr. B. K. Mishra, Principal, Thakur College of Engineering and Technology, for providing the necessary facilities and research infrastructure. We are also thankful to the project coordinators, faculty members, laboratory staff, and our peers for their continued assistance and valuable suggestions."
    p_ack1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ack1.paragraph_format.line_spacing = 1.5
    for run in p_ack1.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    # Student signatures in Acknowledgement
    doc.paragraphs[78].text = "1. Ashmit Singh (Roll No. 10)"
    doc.paragraphs[81].text = "2. Sumit Singh (Roll No. 21)"
    doc.paragraphs[83].text = "3. Shivam Singh (Roll No. 19)"
    doc.paragraphs[86].text = "(Department of Computer Engineering, TCET)"
    for idx in [78, 81, 83, 86]:
        if idx < len(doc.paragraphs):
            for run in doc.paragraphs[idx].runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    # 5. Update Plagiarism Report
    for p_idx in range(88, 95):
        if p_idx < len(doc.paragraphs):
            if "plagiarism" in doc.paragraphs[p_idx].text.lower():
                p_plag = doc.paragraphs[p_idx]
                p_plag.text = "Blue Book Plagiarism Report\n(From Department Turnitin Account Only)"
                p_plag.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p_plag.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(16)
                    run.font.bold = True

    # Clean P94 and any <Times New Roman...> in frontmatter
    for p_idx in range(90, 105):
        if p_idx < len(doc.paragraphs):
            if '<times new roman' in doc.paragraphs[p_idx].text.lower():
                doc.paragraphs[p_idx].text = ''

    # Clean Table 2 headers
    t2 = doc.tables[2]
    for cell in t2.rows[0].cells:
        cell.text = re.sub(r'<[^>]+>', '', cell.text).strip()

    # 6. Update Table 2 (INDEX Page Numbers)
    page_map = {
        'List of Figures': 'I',
        'List of Tables': 'II',
        'Abstract': 'III',
        'Chapter 1': '1',
        '1.1': '2',
        '1.2': '6',
        '1.3': '9',
        '1.4': '12',
        '1.5': '15',
        '1.6': '17',
        'Chapter 2': '19',
        '2.1': '20',
        '2.2': '29',
        '2.3': '36',
        'Chapter 3': '45',
        '3.1': '46',
        '3.2': '54',
        '3.3': '59',
        '3.4': '66',
        '3.5': '72',
        '3.6': '77',
        'Chapter 4': '87',
        '4.1': '88',
        '4.2': '99',
        '4.3': '111',
        '4.4': '121',
        '4.5': '126',
        '4.6': '137',
        '4.7': '145',
        'Chapter 5': '147',
        '5.1': '148',
        '5.2': '157',
        '5.3': '165',
        'Chapter 6': '172',
        '6.1': '173',
        '6.2': '176',
    }

    t2 = doc.tables[2]
    for r in t2.rows:
        col0 = r.cells[0].text.strip()
        col1 = r.cells[1].text.strip()
        key = col0 if col0 in page_map else (col1 if col1 in page_map else None)
        if key and key in page_map:
            r.cells[3].text = page_map[key]
            for p in r.cells[3].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)

    # 7. Update List of Figures (clean out <Times New Roman...> instructions)
    figures = [
        "Figure 1: Escalation of Global Financial Fraud Losses (2018–2026) ........................................ Page 4",
        "Figure 2: The High-Dimensional Fraud Detection Lifecycle ..................................................... Page 7",
        "Figure 3: Ishikawa (Fishbone) Root Cause Diagram of Fraud Detection Failures ..................... Page 33",
        "Figure 4: High-Level Block Diagram of FraudShield AI Proposed Architecture ........................ Page 38",
        "Figure 5: Domain-Specific Model Registry Routing Workflow .................................................. Page 42",
        "Figure 6: Modified CRISP-DM Framework for Real-Time Financial AI ................................... Page 61",
        "Figure 7: Gantt Chart Schedule Across Academic Project Milestones ..................................... Page 75",
        "Figure 8: Use Case Diagram of FraudShield AI Platform ........................................................... Page 79",
        "Figure 9: Unified Structural Class Diagram (Core Model, API, and Data Layers) .................. Page 82",
        "Figure 10: Sequence Diagram: Real-Time Scoring, Explainability, and Graph Lookup ............. Page 85",
        "Figure 11: Data Flow Diagram (DFD Level 0 — Context Diagram) .......................................... Page 90",
        "Figure 12: Data Flow Diagram (DFD Level 1 — Modular Decomposition) ................................ Page 92",
        "Figure 13: Data Flow Diagram (DFD Level 2 — Ingestion, TreeSHAP, NetworkX Graph) ........ Page 94",
        "Figure 14: Microservices and Container Deployment Architecture .......................................... Page 97",
        "Figure 15: Stacking Ensemble Architecture with XGBoost, LightGBM, RF, Meta-Learner ..... Page 100",
        "Figure 16: Algorithmic Process Flow: Dynamic Calibrated Risk Scoring .................................. Page 104",
        "Figure 17: TreeSHAP Feature Attribution Computation Workflow ......................................... Page 107",
        "Figure 18: Graph Intelligence Engine: Cyclic Mule Ring and Community Detection ............... Page 110",
        "Figure 19: UI Snapshot: Enterprise Threat Intelligence Executive Dashboard .......................... Page 113",
        "Figure 20: UI Snapshot: Live Transaction Stream Monitor with Risk Tier Badges ................... Page 115",
        "Figure 21: UI Snapshot: Single Transaction Diagnostic Inspector & Feature Adjuster ............. Page 117",
        "Figure 22: UI Snapshot: TreeSHAP Feature Attribution Waterfall and FCRA Reason Codes ... Page 119",
        "Figure 23: UI Snapshot: Interactive Network Graph & Mule Ring Topology ............................ Page 120",
        "Figure 24: UI Snapshot: MLOps Concept Drift Monitor (KS-Test & PSI Tracking) .................. Page 121",
        "Figure 25: UI Snapshot: Federated Learning Defense Simulator (FedAvg Consensus) .............. Page 122",
        "Figure 26: Confusion Matrix Breakdown across All Four Operational Domains ...................... Page 150",
    ]

    # Clean any font instructions around List of Figures (P105 to P125)
    for p_idx in range(104, 126):
        if p_idx < len(doc.paragraphs):
            txt = doc.paragraphs[p_idx].text.strip()
            if "<times new roman" in txt.lower():
                doc.paragraphs[p_idx].text = ""

    p_fig_start = 108
    for f_idx, fig_text in enumerate(figures):
        p_idx = p_fig_start + f_idx
        if p_idx < len(doc.paragraphs) and p_idx < 120:
            p = doc.paragraphs[p_idx]
            p.text = fig_text
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

    # 8. Update List of Tables
    tables_list = [
        "Table 1: Comprehensive Literature Survey Matrix of Existing Fraud Detection Approaches ....... Page 25",
        "Table 2: Analytical Comparison between Legacy Rule Engines vs. Machine Learning Models ..... Page 31",
        "Table 3: Research Gap Analysis and FraudShield AI Technical Countermeasures ........................ Page 35",
        "Table 4: Functional Requirements Specifications (FR-01 to FR-12) ............................................ Page 48",
        "Table 5: Non-Functional Requirements Specifications (NFR-01 to NFR-08) .............................. Page 52",
        "Table 6: Role-Based Access Control (RBAC) Permission Matrix ................................................. Page 53",
        "Table 7: Technical Feasibility Benchmark and Platform Dependencies ......................................... Page 56",
        "Table 8: Comprehensive Technology Stack Specification .............................................................. Page 68",
        "Table 9: Agile Sprint Backlog, Story Points, and Delivery Milestones ........................................... Page 74",
        "Table 10: Hardware Specification for Training, Staging, and Production Deployment .................. Page 123",
        "Table 11: Benchmark Financial Datasets Summary and Class Distribution Characteristics ............ Page 125",
        "Table 12: Automated Test Suite Execution Matrix (Unit, Integration, and Regression Tests) ....... Page 135",
        "Table 13: Cross-Validation Generalization Gap Analysis (5-Fold CV) ......................................... Page 144",
        "Table 14: Master Empirical Performance Comparison: Default vs. Optimized Model Heads ........ Page 149",
        "Table 15: Optimal Decision Thresholds and Cost Sensitivity per Domain .................................... Page 151",
        "Table 16: TreeSHAP Feature Attribution Summary for High-Risk Transactions ........................... Page 155",
        "Table 17: Federated Learning Aggregation Rounds vs. Multi-Bank Detection Performance .......... Page 157",
        "Table 18: API Latency and Request Throughput Benchmarking Results ........................................ Page 163",
        "Table 19: Comparative Evaluation against Published State-of-the-Art Literature ........................ Page 167",
    ]

    for p_idx in range(120, 142):
        if p_idx < len(doc.paragraphs):
            txt = doc.paragraphs[p_idx].text.strip()
            if "<times new roman" in txt.lower():
                doc.paragraphs[p_idx].text = ""

    p_tbl_start = 126
    for t_idx, tbl_text in enumerate(tables_list):
        p_idx = p_tbl_start + t_idx
        if p_idx < len(doc.paragraphs) and p_idx < 136:
            p = doc.paragraphs[p_idx]
            p.text = tbl_text
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

    # 9. Clean all template guidelines and placeholder text from paragraph 136 onwards!
    # In the template, paragraphs 136+ were dummy guidelines and placeholder Appendix A, B, C!
    # We will delete ALL paragraphs from index 136 up to the end of the template.
    print(f"Removing template guideline paragraphs (from 136 to {len(doc.paragraphs)})...")
    for p in list(doc.paragraphs[136:]):
        delete_paragraph(p)

    # 10. Read the complete markdown source
    with open(markdown_source, "r", encoding="utf-8") as f:
        md_text = f.read()

    # Extract Abstract
    abs_idx = md_text.find("# ABSTRACT")
    abstract_content = ""
    ch1_idx = md_text.find("# Chapter 1: Introduction")
    if ch1_idx == -1:
        ch1_idx = md_text.find("# Chapter 1")

    if abs_idx != -1 and ch1_idx != -1:
        abstract_content = md_text[abs_idx:ch1_idx].strip()

    body_markdown = md_text[ch1_idx:].strip()

    # Add Page Break before ABSTRACT
    doc.add_page_break()

    # Append ABSTRACT
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(24)
    p_abs.paragraph_format.space_after = Pt(18)
    r_abh = p_abs.add_run("ABSTRACT")
    r_abh.font.name = "Times New Roman"
    r_abh.font.size = Pt(18)
    r_abh.font.bold = True

    abs_clean_lines = [l.strip() for l in abstract_content.split("\n") if l.strip() and not l.startswith("#") and not l.startswith("<div") and not l.startswith("---")]
    for aline in abs_clean_lines:
        # Remove any unwanted font tags if present
        aline_clean = re.sub(r'<[^>]+>', '', aline)
        aline_clean = re.sub(r'[*_`]', '', aline_clean)
        p_abt = doc.add_paragraph()
        p_abt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_abt.paragraph_format.line_spacing = 1.5
        p_abt.paragraph_format.space_after = Pt(6)
        r = p_abt.add_run(clean_xml(aline_clean))
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    doc.add_page_break()

    # 11. Parse and append full Chapters 1 to 6, References, Research Paper, and Real Appendices
    print("Appending chapters, tables, algorithms, and figures...")
    lines = body_markdown.split("\n")
    i = 0
    in_code_block = False
    code_lines = []

    def insert_diagram(image_filename, caption_text):
        img_path = os.path.join(figures_dir, image_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(5.8))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(14)
            p_cap.paragraph_format.keep_with_next = True
            run_cap = p_cap.add_run(caption_text)
            run_cap.font.name = "Times New Roman"
            run_cap.font.size = Pt(10.5)
            run_cap.font.bold = True
            run_cap.font.italic = True

    while i < len(lines):
        line = lines[i]

        if '<div style="page-break-before: always;"></div>' in line:
            doc.add_page_break()
            i += 1
            continue

        if line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F3F4F6")
                set_cell_margins(cell, 150, 150, 200, 200)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(clean_xml(code_text))
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(31, 41, 55)
                code_lines = []
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
                data_start = 1
                if len(table_lines) > 1 and re.match(r'^\|?[\s\-:|]+\|?$', table_lines[1]):
                    data_start = 2

                rows_data = []
                for tline in table_lines[data_start:]:
                    cols = [c.strip() for c in tline.strip("|").split("|")]
                    rows_data.append(cols)

                if headers and rows_data:
                    num_cols = len(headers)
                    table = doc.add_table(rows=len(rows_data) + 1, cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, "D1D5DB", sz="4")

                    hdr_cells = table.rows[0].cells
                    for col_idx, htext in enumerate(headers):
                        if col_idx < num_cols:
                            cell = hdr_cells[col_idx]
                            set_cell_background(cell, "E5E7EB")
                            set_cell_margins(cell, 120, 120, 140, 140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_after = Pt(2)
                            h_clean = re.sub(r'<[^>]+>', '', htext)
                            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_clean)))
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10.5)
                            run.font.bold = True

                    for r_idx, r_data in enumerate(rows_data):
                        row_cells = table.rows[r_idx + 1].cells
                        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFB"
                        for col_idx in range(num_cols):
                            cell = row_cells[col_idx]
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, 100, 100, 120, 120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_after = Pt(2)
                            val = r_data[col_idx] if col_idx < len(r_data) else ""
                            val_clean = re.sub(r'<[^>]+>', '', val)
                            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', val_clean)))
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)

                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if line.startswith("# "):
            h_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            if "Chapter" in h_text or "CERTIFICATE" in h_text or "INDEX" in h_text or "ABSTRACT" in h_text or "REFERENCES" in h_text or "RESEARCH PAPER" in h_text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h_clean = re.sub(r'<[^>]+>', '', h_text)
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_clean)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
            run.font.bold = True
            i += 1
            continue

        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            h_clean = re.sub(r'<[^>]+>', '', h_text)
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_clean)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.font.bold = True

            # If this is Section 2.3.2 High-Level System Architecture, insert Figure 4
            if "2.3.2" in h_text or "High-Level System Architecture" in h_text:
                insert_diagram("fig4_architecture.jpg", "Figure 4: High-Level Block Diagram of FraudShield AI Proposed Architecture")

            i += 1
            continue

        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            h_clean = re.sub(r'<[^>]+>', '', h_text)
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_clean)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.font.bold = True

            # Check if this heading matches other diagrams
            if "4.2.2" in h_text or "Hybrid Stacking Ensemble" in h_text:
                insert_diagram("fig15_stacking.jpg", "Figure 15: Stacking Ensemble Architecture with XGBoost, LightGBM, RF, and Meta-Learner")
            elif "4.2.4" in h_text or "TreeSHAP Feature Attribution" in h_text:
                insert_diagram("fig17_treeshap.jpg", "Figure 17: TreeSHAP Feature Attribution Computation Workflow & FCRA Mapping")
            elif "4.2.5" in h_text or "Cyclic Mule Ring" in h_text:
                insert_diagram("fig18_mule_ring.jpg", "Figure 18: Graph Intelligence Engine: Cyclic Mule Ring and Community Detection")

            i += 1
            continue

        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            h_clean = re.sub(r'<[^>]+>', '', h_text)
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_clean)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            i += 1
            continue

        if line.strip() in ["---", "___", "***"]:
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # Clean any stray <...> font instructions from body paragraphs
        cleaned_line = clean_xml(line.strip())
        if re.match(r'^<Times New Roman.*?>$', cleaned_line, re.IGNORECASE):
            i += 1
            continue

        # Replace roll numbers if older variants exist
        cleaned_line = cleaned_line.replace("Roll No. 101", "Roll No. 10")
        cleaned_line = cleaned_line.replace("Roll No. 102", "Roll No. 21")
        cleaned_line = cleaned_line.replace("Roll No. 103", "Roll No. 19")

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)

        if cleaned_line.startswith("- ") or cleaned_line.startswith("* "):
            p.paragraph_format.left_indent = Inches(0.3)
            bullet_text = cleaned_line[2:]
            run = p.add_run("• " + re.sub(r'[*_`]', '', bullet_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        elif re.match(r'^\d+\.\s', cleaned_line):
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(re.sub(r'[*_`]', '', cleaned_line))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(re.sub(r'[*_`]', '', cleaned_line))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        i += 1

    print(f"Saving final cleaned and formatted document to: {output_path}...")
    doc.save(output_path)
    print("SUCCESS! Completed building FraudShield_AI_Final_Blue_Book.docx")

if __name__ == "__main__":
    main()
