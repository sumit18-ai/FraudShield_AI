import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def clean_xml(text):
    if not text:
        return ""
    # Filter out invalid XML characters
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', str(text))

def build_docx():
    docs_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    md_path = os.path.join(docs_dir, "FraudShield_AI_Blue_Book_Report.md")
    docx_path = os.path.join(docs_dir, "FraudShield_AI_Blue_Book_Report.docx")

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Configure Margins: 1.25" Left, 1.0" Right, Top, Bottom
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Base style configurations
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(17, 24, 39)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    lines = md_text.split("\n")
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Check for page break
        if '<div style="page-break-before: always;"></div>' in line:
            doc.add_page_break()
            i += 1
            continue

        # Check for code blocks
        if line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                # Write code block
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F3F4F6")
                set_cell_margins(cell, 150, 150, 200, 200)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(clean_xml(code_text))
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(31, 41, 55)
                code_lines = []
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
                data_start = 1
                if len(table_lines) > 1 and re.match(r'^\|?[\s\-:|]+\|?$', table_lines[1]):
                    data_start = 2

                rows_data = []
                for tline in table_lines[data_start:]:
                    cols = [c.strip() for c in tline.strip("|").split("|")]
                    rows_data.append(cols)

                if headers and rows_data:
                    num_cols = len(headers)
                    table = doc.add_table(rows=len(rows_data) + 1, cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, "D1D5DB", sz="4")

                    # Add headers
                    hdr_cells = table.rows[0].cells
                    for col_idx, htext in enumerate(headers):
                        if col_idx < num_cols:
                            cell = hdr_cells[col_idx]
                            set_cell_background(cell, "E5E7EB")
                            set_cell_margins(cell, 120, 120, 140, 140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_after = Pt(2)
                            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', htext)))
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10.5)
                            run.font.bold = True

                    # Add data rows
                    for r_idx, r_data in enumerate(rows_data):
                        row_cells = table.rows[r_idx + 1].cells
                        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFB"
                        for col_idx in range(num_cols):
                            cell = row_cells[col_idx]
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, 100, 100, 120, 120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_after = Pt(2)
                            val = r_data[col_idx] if col_idx < len(r_data) else ""
                            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', val)))
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)

                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_after = Pt(6)
            continue

        # Markdown Headings
        if line.startswith("# "):
            h_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            if "Chapter" in h_text or "CERTIFICATE" in h_text or "INDEX" in h_text or "ABSTRACT" in h_text or "REFERENCES" in h_text or "Project Report" in h_text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_text)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
            run.font.bold = True
            i += 1
            continue

        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_text)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.font.bold = True
            i += 1
            continue

        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_text)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.font.bold = True
            i += 1
            continue

        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml(re.sub(r'[*_`]', '', h_text)))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            i += 1
            continue

        # Horizontal rules
        if line.strip() in ["---", "___", "***"]:
            i += 1
            continue

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Normal paragraphs and lists
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)

        cleaned_line = clean_xml(line.strip())
        if cleaned_line.startswith("- ") or cleaned_line.startswith("* "):
            p.paragraph_format.left_indent = Inches(0.3)
            bullet_text = cleaned_line[2:]
            run = p.add_run("• " + re.sub(r'[*_`]', '', bullet_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        elif re.match(r'^\d+\.\s', cleaned_line):
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(re.sub(r'[*_`]', '', cleaned_line))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(re.sub(r'[*_`]', '', cleaned_line))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX report at: {docx_path}")

if __name__ == "__main__":
    build_docx()
