import os
import docx

template_path = r'c:\Users\Sumit\OneDrive\Desktop\FraudShield_AI\FraudShield_AI\docs\Draft Blue book format for 2026-27 -Final.docx'
doc = docx.Document(template_path)

print(f"Total sections: {len(doc.sections)}")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Dump all paragraphs with text and formatting info
with open("template_structure.txt", "w", encoding="utf-8") as f:
    f.write(f"=== SECTIONS ({len(doc.sections)}) ===\n")
    for s_idx, sec in enumerate(doc.sections):
        f.write(f"Section {s_idx}: Top={sec.top_margin.inches:.2f}, Bottom={sec.bottom_margin.inches:.2f}, Left={sec.left_margin.inches:.2f}, Right={sec.right_margin.inches:.2f}\n")
        f.write(f"  Header runs: {len(sec.header.paragraphs[0].runs) if sec.header.paragraphs else 0}\n")
        
    f.write(f"\n=== PARAGRAPHS ({len(doc.paragraphs)}) ===\n")
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        align = p.alignment
        runs_info = []
        for r in p.runs:
            fn = r.font.name
            fs = r.font.size.pt if r.font.size else None
            b = r.bold
            it = r.italic
            runs_info.append(f"{fn}/{fs}pt/B={b}/I={it}")
        f.write(f"P{p_idx:3d} [{p.style.name}] (align={align}) runs=[{', '.join(runs_info)}]: {p_text}\n")
        
    f.write(f"\n=== TABLES ({len(doc.tables)}) ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} cols\n")
        for r_idx, row in enumerate(table.rows):
            f.write(f"  R{r_idx}: {[c.text.strip().replace(chr(10), ' ') for c in row.cells]}\n")

print("Saved detailed analysis to template_structure.txt")
